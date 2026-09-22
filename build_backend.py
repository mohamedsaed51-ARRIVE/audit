"""Builds the Backend workbook (headers only) that becomes the Google Sheet database.
No formulas: Apps Script (Code.gs) computes everything on write, since Google Sheets
formulas are not portable the same way and the write path must be server-side anyway."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(__file__)
schema = json.load(open(os.path.join(HERE, "schema.json"), encoding="utf-8"))

AR_FONT = "Arial"
NAVY = "1B2A4A"
WHITE = "FFFFFF"

wb = Workbook()
wb.remove(wb.active)

def header_row(ws, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=AR_FONT, bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", readingOrder=2)
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True

sys_cols = schema["system_columns"]

for ct in schema["control_types"]:
    ws = wb.create_sheet(ct["sheet"])
    headers = [c["label"] for c in sys_cols[:4]]  # record_id, entry_date, entry_time, entered_by
    headers += [f["label"] for f in ct["fields"]]
    headers += [c["label"] for c in ct.get("computed", [])]
    headers += [c["label"] for c in sys_cols[4:]]  # last_modified, modified_by
    header_row(ws, headers)
    for i in range(len(headers)):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A"].width = 18

# ---- الإعدادات ----
ws = wb.create_sheet("الإعدادات")
ws.sheet_view.rightToLeft = True
row = 1
ws.cell(row=row, column=1, value="نوع القائمة").font = Font(bold=True, name=AR_FONT)
ws.cell(row=row, column=2, value="القيمة").font = Font(bold=True, name=AR_FONT)
ws.cell(row=row, column=3, value="المعرّف (للفروع فقط)").font = Font(bold=True, name=AR_FONT)
ws.cell(row=row, column=4, value="الحالة (للفروع فقط)").font = Font(bold=True, name=AR_FONT)
row += 1
for key, values in schema["lists"].items():
    for i, v in enumerate(values, start=1):
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=v)
        if key == "Branches":
            # Seed stable Branch IDs (BR-001, BR-002, ...) and default status "نشط" at build time,
            # so a freshly generated workbook already supports Branch Management from day one.
            ws.cell(row=row, column=3, value="BR-%03d" % i)
            ws.cell(row=row, column=4, value="نشط")
        row += 1
row += 2
ws.cell(row=row, column=1, value="الأهداف الرقابية").font = Font(bold=True, name=AR_FONT, color=WHITE)
ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=NAVY)
row += 1
for key, val in schema["targets"].items():
    ws.cell(row=row, column=1, value=key)
    ws.cell(row=row, column=2, value=val)
    row += 1
row += 2
ws.cell(row=row, column=1, value="فئات عمر الحالات المتأخرة").font = Font(bold=True, name=AR_FONT, color=WHITE)
ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=NAVY)
row += 1
ws.cell(row=row, column=1, value="من"); ws.cell(row=row, column=2, value="إلى"); ws.cell(row=row, column=3, value="التصنيف")
row += 1
for a, b, label in schema["age_buckets"]:
    ws.cell(row=row, column=1, value=a)
    ws.cell(row=row, column=2, value=b)
    ws.cell(row=row, column=3, value=label)
    row += 1

# ---- عداد التسلسل (Sequence) — يستخدمه Apps Script لتوليد أرقام السجلات دون تكرار ----
ws = wb.create_sheet("_Sequence")
ws.sheet_view.rightToLeft = True
ws.cell(row=1, column=1, value="نوع الرقابة")
ws.cell(row=1, column=2, value="آخر رقم مستخدم")
r = 2
for ct in schema["control_types"]:
    ws.cell(row=r, column=1, value=ct["key"])
    ws.cell(row=r, column=2, value=0)
    r += 1

# ---- الموظفون + _AuditLog — تُنشأ فارغة (بلا أي حساب افتراضي)؛ أول حساب Administrator يُنشأ
# لاحقًا من واجهة Dashboard عبر إجراء bootstrapAdmin (يعمل مرة واحدة فقط طالما التاب فارغ) ----
for admin_sheet in schema.get("admin_sheets", []):
    ws = wb.create_sheet(admin_sheet["sheet"])
    ws.sheet_view.rightToLeft = True
    header_row(ws, admin_sheet["headers"])
    for i in range(len(admin_sheet["headers"])):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A"].width = 18

out = os.path.join(HERE, "ARRIVE_Backend.xlsx")
wb.save(out)
print("saved", out, "sheets:", wb.sheetnames)
